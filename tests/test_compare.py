"""Comparison Excel: a summary row per experiment (lr, epoch, sequence, method, mAP / AP50 / AP75 as the plain and the
frame-weighted mean) and, when the experiments have several sequences, one sheet per experiment with a row per
sequence and the two means."""
import json
import tempfile
from pathlib import Path

import openpyxl

from stark_ft.test.compare import export_comparison

HEADER = ("experiment", "lr", "epoch", "sequence", "method", "mAP", "AP50", "AP75")
SUMMARY = (*HEADER, "mAP weighted", "AP50 weighted", "AP75 weighted")


def _experiment(outputs, name, config, seq_maps, earlier=False, frames=None):
    """seq_maps: {sequence: mAP}; AP50 = mAP + 0.1, AP75 = mAP - 0.1. frames: {sequence: number of frames}."""
    rows = []
    for seq, m in seq_maps.items():
        row = {"sequence": seq, "mAP": m, "AP50": m + 0.1, "AP75": m - 0.1, "F_opt": 0.5}
        if earlier:  # metrics.json of earlier versions: these values as legacy_*, other AP values under the plain names
            row = {"sequence": seq, "AP": 0.9, "AP50": 0.9, "AP75": 0.9, "F_opt": 0.5,
                   "legacy_AP": m, "legacy_AP50": m + 0.1, "legacy_AP75": m - 0.1}
        if frames:
            row["frames"] = frames[seq]
        rows.append(row)
    mean = {k: v for k, v in rows[0].items() if k != "sequence"}
    d = outputs / name / "metrics"
    d.mkdir(parents=True)
    (d / "metrics.json").write_text(json.dumps({
        "summary": {"mean_over_sequences": mean, "pooled": mean, "n_sequences": len(rows)},
        "config": config, "per_sequence": rows}))


def _rows(sheet):
    return [tuple(round(v, 6) if isinstance(v, float) else v for v in r) for r in sheet.iter_rows(values_only=True)]


def test_one_sequence_summary_only():
    with tempfile.TemporaryDirectory() as tmp:
        outputs = Path(tmp)
        _experiment(outputs, "b_exp", {"ft_mode": "online", "ft_epochs_online": 15}, {"bull": 0.5},
                    frames={"bull": 2000})
        _experiment(outputs, "a_exp", {"ft_mode": "init", "ft_lr": 1e-5, "update_mode": "max"}, {"bull": 0.3},
                    earlier=True, frames={"bull": 2000})
        _experiment(outputs, "c_exp", {"ft_mode": "none", "update_interval": 99999}, {"bull": 0.7})  # no frames
        wb = openpyxl.load_workbook(export_comparison(["b_exp", "a_exp", "c_exp"], outputs / "c.xlsx", outputs))
        assert wb.sheetnames == ["summary"]
        assert _rows(wb["summary"]) == [
            SUMMARY,
            ("b_exp", "0.0001", "15+15", "bull", "orj", 0.5, 0.6, 0.4, 0.5, 0.6, 0.4),  # the given order
            ("a_exp", "1e-05", "15+0", "bull", "max", 0.3, 0.4, 0.2, 0.3, 0.4, 0.2),
            ("c_exp", "-", "0+0", "bull", "noupd", 0.7, 0.8, 0.6, None, None, None)]  # weighted: unknown


def test_several_sequences_one_sheet_per_experiment():
    with tempfile.TemporaryDirectory() as tmp:
        outputs = Path(tmp)
        frames = {"ballet": 100, "bull": 300, "car1": 600}
        _experiment(outputs, "exp1", {}, {"ballet": 0.2, "bull": 0.5, "car1": 0.8}, frames=frames)
        _experiment(outputs, "exp2", {"ft_mode": "none"}, {"ballet": 0.1, "bull": 0.4, "car1": 0.4}, frames=frames)
        wb = openpyxl.load_workbook(export_comparison(["exp1", "exp2"], outputs / "c.xlsx", outputs))
        assert wb.sheetnames == ["summary", "1 orj lr0.0001 15+1", "2 orj 0+0"]
        # weighted: (0.2 * 100 + 0.5 * 300 + 0.8 * 600) / 1000 = 0.65; (0.01 + 0.12 + 0.24) / 1 = 0.37
        assert _rows(wb["summary"])[1:] == [
            ("exp1", "0.0001", "15+1", "3 sequences", "orj", 0.5, 0.6, 0.4, 0.65, 0.75, 0.55),
            ("exp2", "-", "0+0", "3 sequences", "orj", 0.3, 0.4, 0.2, 0.37, 0.47, 0.27)]
        sheet = _rows(wb["2 orj 0+0"])
        assert sheet[0] == HEADER and [r[3] for r in sheet[1:]] == ["ballet", "bull", "car1", "mean", "weighted mean"]
        assert sheet[-2] == (None, None, None, "mean", None, 0.3, 0.4, 0.2)  # plain mean
        assert sheet[-1] == (None, None, None, "weighted mean", None, 0.37, 0.47, 0.27)  # weighted by frames
        assert wb["2 orj 0+0"]["F5"].font.bold and wb["2 orj 0+0"]["F6"].font.bold
        assert not wb["2 orj 0+0"]["F4"].font.bold


if __name__ == "__main__":  # without pytest: python -m tests.test_compare
    for _name, _fn in sorted(globals().items()):
        if _name.startswith("test_") and callable(_fn):
            _fn()
            print("PASS", _name)
